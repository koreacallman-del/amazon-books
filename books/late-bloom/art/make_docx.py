"""워드 파일 생성 — 「시계가 멈춘 마을」

  python3 make_docx.py

그림은 페이지 배경(글 뒤)으로 깔고, 글은 편집 가능한 텍스트로 올린다.
저자가 워드에서 문장을 직접 고칠 수 있어야 하므로 글자를 그림에 굽지 않는다.

산출물: out/시계가-멈춘-마을.docx  (8.5×8.5in, 133쪽)
"""

import io
import os

import pymupdf
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ink import to_pdf, OUTER
from pages import P
from build import compose, load_texts, OUT

DPI = 150                 # 검토용. 인쇄는 late-bloom-interior.pdf를 쓴다
TRIM = 8.5
GUTTER_IN = 0.375         # 책등 쪽 안전 여백
OUTER_IN = 0.55           # 바깥 여백

NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def render_images():
    """글 없이 그림만 렌더 → 페이지별 PNG 바이트."""
    texts = load_texts()
    pdf_pages = []
    for n in range(1, 134):
        bleed, content, _ = compose(n, texts)
        pdf_pages.append((bleed + content, []))       # 글은 넣지 않는다
    tmp = os.path.join(OUT, "_art_only.pdf")
    to_pdf(pdf_pages, tmp)

    doc = pymupdf.open(tmp)
    mat = pymupdf.Matrix(DPI / 72.0, DPI / 72.0)
    images = [p.get_pixmap(matrix=mat, alpha=False).tobytes("png")
              for p in doc]
    doc.close()
    os.remove(tmp)
    return images, texts


def send_behind(run, w_emu, h_emu):
    """인라인 그림을 페이지 원점에 고정된 '글 뒤' 배경으로 바꾼다."""
    inline = run._r.find(qn("w:drawing"))[0]
    anchor = OxmlElement("wp:anchor")
    for k, v in (("distT", "0"), ("distB", "0"), ("distL", "0"),
                 ("distR", "0"), ("simplePos", "0"), ("relativeHeight", "1"),
                 ("behindDoc", "1"), ("locked", "0"), ("layoutInCell", "1"),
                 ("allowOverlap", "1")):
        anchor.set(k, v)

    sp = OxmlElement("wp:simplePos"); sp.set("x", "0"); sp.set("y", "0")
    anchor.append(sp)
    for tag, off in (("wp:positionH", "0"), ("wp:positionV", "0")):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", "page")
        o = OxmlElement("wp:posOffset"); o.text = off
        pos.append(o)
        anchor.append(pos)

    ext = OxmlElement("wp:extent")
    ext.set("cx", str(int(w_emu))); ext.set("cy", str(int(h_emu)))
    anchor.append(ext)
    ee = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        ee.set(k, "0")
    anchor.append(ee)
    anchor.append(OxmlElement("wp:wrapNone"))

    for tag in ("wp:docPr", "wp:cNvGraphicFramePr"):
        el = inline.find(qn(tag))
        if el is not None:
            anchor.append(el)
    graphic = inline.find(
        qn("a:graphic")) if inline.find(qn("a:graphic")) is not None \
        else inline[-1]
    anchor.append(graphic)

    parent = inline.getparent()
    parent.remove(inline)
    parent.append(anchor)


def text_top_in(n, lines):
    """SVG 조판과 같은 세로 위치를 인치로 환산."""
    _, tpos = P[n]
    lead = 6.6 if n == 1 else 6.0
    if tpos == "top":
        y0 = 24
    elif tpos == "mid":
        y0 = 48 - (len(lines) - 1) * lead / 2
    else:
        y0 = 100 - OUTER - 6 - (len(lines) - 1) * lead
    return max(0.35, (y0 / 100.0) * TRIM - 0.30)


def build_docx():
    images, texts = render_images()
    d = Document()

    sec = d.sections[0]
    sec.page_width = Inches(TRIM)
    sec.page_height = Inches(TRIM)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Inches(0))
    sec.header_distance = Inches(0)
    sec.footer_distance = Inches(0)

    st = d.styles["Normal"]
    st.font.name = "바탕"
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)

    emu = int(TRIM * 914400)

    for n in range(1, 134):
        lines = texts.get(n, [])
        _, tpos = P[n]

        # 배경 그림
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)
        run = p.add_run()
        run.add_picture(io.BytesIO(images[n - 1]), width=Inches(TRIM))
        send_behind(run, emu, emu)

        # 글 — 그림 위에 올라간다
        if lines:
            top = text_top_in(n, lines)
            for i, s in enumerate(lines):
                tp = d.add_paragraph()
                tp.paragraph_format.space_before = (
                    Inches(top) if i == 0 else Pt(0))
                tp.paragraph_format.space_after = Pt(0)
                tp.paragraph_format.line_spacing = Pt(20 if n == 1 else 18)
                if n == 1 or tpos == "mid":
                    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    tp.paragraph_format.left_indent = Inches(0.9)
                    tp.paragraph_format.right_indent = Inches(0.9)
                else:
                    tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    inner, outer = Inches(GUTTER_IN + 0.6), Inches(OUTER_IN)
                    if n % 2:                       # 홀수 = 책등 왼쪽
                        tp.paragraph_format.left_indent = inner
                        tp.paragraph_format.right_indent = outer
                    else:
                        tp.paragraph_format.left_indent = outer
                        tp.paragraph_format.right_indent = inner
                r = tp.add_run(s)
                r.font.name = "바탕"
                rpr = r._element.get_or_add_rPr()
                rf = rpr.find(qn("w:rFonts"))
                if rf is None:
                    rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
                for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
                    rf.set(qn(a), "바탕")
                if n == 1:
                    r.font.size = Pt(24 if i == 0 else 13)
                    r.font.color.rgb = (RGBColor(0x1C, 0x1A, 0x17) if i == 0
                                        else RGBColor(0x56, 0x50, 0x4A))
                else:
                    r.font.size = Pt(12)
                    r.font.color.rgb = RGBColor(0x1C, 0x1A, 0x17)

        if n < 133:
            d.add_page_break()

    path = os.path.join(OUT, "시계가-멈춘-마을.docx")
    d.save(path)
    return path


if __name__ == "__main__":
    path = build_docx()
    mb = os.path.getsize(path) / 1024 / 1024
    print(f"완료 → {path}  ({mb:.1f} MB · 133쪽 · {DPI}dpi)")
